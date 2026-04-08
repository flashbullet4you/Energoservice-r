"""Обработка файлов: конвертация DWG, извлечение текста, разбиение на чанки."""

import hashlib
import os
import subprocess
from pathlib import Path
from typing import List

import openpyxl
import pypdf
from docx import Document

from config import (
    DOC_DIR,
    DWG_CACHE_DIR,
    DWG_CONVERTER_CMD,
    CHUNK_SIZE,
    CHUNK_OVERLAP,
    logger,
)


# ==================== КОНВЕРТАЦИЯ DWG ====================
def convert_dwgs():
    """Пакетная конвертация всех DWG в директории -> PDF."""
    os.makedirs(DWG_CACHE_DIR, exist_ok=True)
    dwg_files = list(Path(DOC_DIR).rglob("*.dwg"))
    if not dwg_files:
        return

    logger.info(f"📐 Найдено {len(dwg_files)} DWG файлов. Запуск конвертации...")
    cmd = [
        DWG_CONVERTER_CMD,
        "-i",
        DOC_DIR,
        "-o",
        DWG_CACHE_DIR,
        "-v",
        "ACAD2018",
        "-f",
        "PDF",
        "-r",
        "-x",
        "-l",
    ]
    try:
        subprocess.run(cmd, check=True, timeout=600, capture_output=True)
        logger.info("✅ DWG -> PDF конвертация завершена.")
    except subprocess.CalledProcessError as e:
        logger.warning(f"⚠️ Ошибка конвертера: {e.stderr.decode()[:200]}")
    except FileNotFoundError:
        logger.warning(
            f"⚠️ Конвертер '{DWG_CONVERTER_CMD}' не найден в PATH. DWG будут пропущены."
        )
    except Exception as e:
        logger.warning(f"⚠️ Конвертация DWG пропущена: {e}")


# ==================== ИЗВЛЕЧЕНИЕ ТЕКСТА ====================
def extract_text(file_path: str) -> str:
    ext = Path(file_path).suffix.lower()
    try:
        if ext == ".pdf":
            return "\n".join(
                page.extract_text() or "" for page in pypdf.PdfReader(file_path).pages
            )
        elif ext == ".docx":
            return "\n".join(p.text for p in Document(file_path).paragraphs)
        elif ext == ".xlsx":
            return "\n".join(
                " ".join(str(c) for c in row if c)
                for sheet in openpyxl.load_workbook(file_path, data_only=True)
                for row in sheet.iter_rows(values_only=True)
            )
        return ""
    except Exception as e:
        logger.error(f"Ошибка чтения {file_path}: {e}")
        return ""


# ==================== РАЗБИЕНИЕ НА ЧАНКИ ====================
def chunk_text(text: str, source: str) -> List[dict]:
    if not text.strip():
        return []
    words, chunks, i = text.split(), [], 0
    while i < len(words):
        end = min(i + CHUNK_SIZE, len(words))
        chunk = " ".join(words[i:end])
        chunks.append(
            {
                "text": chunk,
                "source": source,
                "id": f"{source}_{hashlib.md5(chunk.encode()).hexdigest()[:8]}",
            }
        )
        i += CHUNK_SIZE - CHUNK_OVERLAP
    return chunks
